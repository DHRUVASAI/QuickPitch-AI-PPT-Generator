import sys
import os
from pptx import Presentation

pptx_path = os.path.join(os.getcwd(), 'scratch', 'real_pipeline_hindi_deck.pptx')
if not os.path.exists(pptx_path):
    print(f"File not found: {pptx_path}")
    sys.exit(1)

prs = Presentation(pptx_path)
print(f"Total Slides in Deck: {len(prs.slides)}\n")

for i, slide in enumerate(prs.slides, 1):
    print(f"=== SLIDE {i} ===")
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t and t not in texts:
                    texts.append(t)
    
    for t in texts:
        print(f"  • {t}")
    
    # Speaker notes
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            print(f"  [Speaker Notes]: {notes}")
    print()
